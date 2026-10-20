import json
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

source = Path(sys.argv[1])
outdir = Path(sys.argv[2])
outdir.mkdir(parents=True, exist_ok=True)
data = json.loads(source.read_text(encoding="utf-8"))
out = outdir / "CyberLab-Studio-Pro-v1.0.0.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
BG = RGBColor(7, 11, 18)
SURF = RGBColor(15, 23, 36)
SURF2 = RGBColor(20, 40, 75)
FG = RGBColor(245, 247, 251)
MUTED = RGBColor(168, 179, 197)
BLUE = RGBColor(114, 167, 255)
CYAN = RGBColor(95, 214, 168)
ORANGE = RGBColor(255, 181, 107)
RED = RGBColor(255, 138, 128)
BORDER = RGBColor(38, 50, 70)

def add_bg(slide):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)

def text(slide, x, y, w, h, content, size=18, color=FG, bold=False, font="Aptos", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    p = frame.paragraphs[0]
    p.text = content
    p.alignment = align
    r = p.runs[0]
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box

def title(slide, heading, subtitle=None):
    text(slide, 0.7, 0.45, 12, 0.55, heading, 25, FG, True)
    if subtitle:
        text(slide, 0.72, 1.03, 11.8, 0.42, subtitle, 10.5, MUTED)
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.42), Inches(1.1), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()

def footer(slide, number):
    text(slide, 0.72, 7.15, 6, 0.18, "CYBERLAB STUDIO PRO  •  v1.0.0", 7, MUTED, True)
    text(slide, 12.2, 7.15, 0.4, 0.18, str(number), 7, MUTED, True, align=PP_ALIGN.RIGHT)

def card(slide, x, y, w, h, heading, body="", accent=BLUE, stat=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = SURF
    shape.line.color.rgb = BORDER
    shape.line.width = Pt(1)
    mark = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.18), Inches(y + 0.2), Inches(0.08), Inches(0.5))
    mark.fill.solid()
    mark.fill.fore_color.rgb = accent
    mark.line.fill.background()
    text(slide, x + 0.38, y + 0.17, w - 0.55, 0.35, heading, 12, FG, True)
    if stat is not None:
        text(slide, x + 0.38, y + 0.7, w - 0.5, 0.6, str(stat), 28, accent, True)
    if body:
        text(slide, x + 0.38, y + (1.28 if stat is not None else 0.65), w - 0.55, h - (1.42 if stat is not None else 0.8), body, 9.5, MUTED)

def bullets(slide, x, y, w, h, items, size=14, color=FG):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = "•  " + item
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(8)

def pill(slide, x, y, w, label, fill=SURF2, color=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.42))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    text(slide, x, y + 0.075, w, 0.22, label, 9, color, True, align=PP_ALIGN.CENTER)

def base(heading, subtitle=None):
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    title(slide, heading, subtitle)
    footer(slide, len(prs.slides))
    return slide

slide = prs.slides.add_slide(blank)
add_bg(slide)
text(slide, 0.75, 0.75, 3, 0.3, "CYBERLAB", 12, BLUE, True)
text(slide, 0.75, 1.25, 11.7, 0.8, "Studio Pro", 34, FG, True)
text(slide, 0.75, 2.08, 11.5, 0.65, "Complete Lab Platform", 25, MUTED, True)
text(slide, 0.75, 2.84, 5, 0.45, "Release v1.0.0", 15, CYAN, True)
for i, (stat, label, color) in enumerate([("30", "versioned labs", BLUE), ("120", "guided tasks", CYAN), ("90", "progressive hints", ORANGE)]):
    card(slide, 0.75 + i * 2.3, 4.05, 2.05, 1.55, label, accent=color, stat=stat)
text(slide, 0.75, 6.35, 11.4, 0.5, "Hands-on cloud-security training with isolated learner state, real validation workflows, and role-based administration.", 13, FG)
footer(slide, 1)

slide = base("Platform overview", "One application, six learning domains, isolated state per learner")
items = [
    ("Cloud fundamentals", "Labs 01–02", "shared responsibility + IAM reasoning", BLUE),
    ("AWS-compatible", "03–18", "IAM, VPC, EC2, storage, exposure", CYAN),
    ("Azure-compatible", "19–22", "Entra, RBAC, network, storage, VM", BLUE),
    ("Sentinel", "23–24", "SIEM, detection, incident workflows", ORANGE),
    ("KQL", "25–26", "query language + investigations", CYAN),
    ("Defender XDR", "27–30", "incident correlation + response", RED),
]
for i, (area, labs, body, color) in enumerate(items):
    card(slide, 0.75 + (i % 3) * 4.1, 1.75 + (i // 3) * 2.15, 3.7, 1.78, f"{area}  ·  {labs}", body, accent=color)

slide = base("Architecture", "Server-controlled trust boundaries; no host-shell execution")
for i, (heading, body) in enumerate([("Browser", "Next.js UI"), ("Auth", "Firebase + HttpOnly session"), ("Policy", "RBAC + quotas + rate limits"), ("Lab engine", "Versioned per-user state"), ("SQLite", "Prisma source of truth")]):
    x = 0.55 + i * 2.55
    card(slide, x, 2.1, 2.22, 1.45, heading, body, accent=BLUE if i < 3 else CYAN)
    if i < 4:
        text(slide, x + 2.23, 2.72, 0.28, 0.25, "→", 18, MUTED, True, align=PP_ALIGN.CENTER)
bullets(slide, 0.8, 4.25, 11.8, 1.8, [
    "AWS/Azure-style commands mutate only the authenticated learner environment.",
    "KQL, Sentinel, and Defender workflows share persisted security datasets and event IDs.",
    "Completion, flags, achievements, and privileged administration are server-authorized.",
    "Production AWS/Azure accounts are not contacted by default.",
], 12.5)

slide = base("Identity & account lifecycle", "Admin-controlled onboarding with optional TOTP MFA")
card(slide, 0.75, 1.75, 3.7, 2.0, "Administrator", "First-run setup token creates the initial admin. Existing admins provision instructor accounts and manage users.", accent=RED)
card(slide, 4.8, 1.75, 3.7, 2.0, "Instructor", "Creates seat-limited, expiring enrollment links. Each consumed seat becomes an assigned student account.", accent=ORANGE)
card(slide, 8.85, 1.75, 3.7, 2.0, "Student", "Opens the instructor link and chooses a username/password. Public student self-signup is disabled.", accent=CYAN)
bullets(slide, 0.9, 4.25, 11.4, 1.8, [
    "Username/email login with persistent attempt cooldown; Firebase abuse controls remain additional protection.",
    "Google sign-in and email/password authentication are configured through Firebase.",
    "Firebase TOTP 2FA enrollment uses a locally rendered QR; sign-in resolves the real second-factor challenge.",
    "Suspension, role checks, and session authorization are enforced server-side.",
], 11.8)

slide = base("AWS-compatible lab environment", "Stateful command workflows with explicit unsupported-operation errors")
for i, (heading, body, color) in enumerate([
    ("Identity", "STS, IAM users/groups/roles, policy evaluation, access keys", BLUE),
    ("Network", "VPC, subnets, routes, security groups, NACLs, flow/audit evidence", CYAN),
    ("Compute", "EC2 state, metadata, patch posture, tags, role relationships", ORANGE),
    ("Storage", "S3 buckets/objects, permissions, encryption, EBS/snapshots", RED),
]):
    card(slide, 0.75 + (i % 2) * 6.0, 1.75 + (i // 2) * 2.15, 5.55, 1.8, heading, body, accent=color)
text(slide, 0.9, 6.25, 11.7, 0.45, "Commands are parsed and dispatched inside CyberLab; no learner command is sent to the host shell.", 11.5, MUTED, True)

slide = base("Azure-compatible lab environment", "Azure CLI-style workflows over isolated learner state")
for i, (heading, body, color) in enumerate([
    ("Entra identity", "Users, groups, membership, service identities", BLUE),
    ("RBAC", "Role definitions, assignments, scopes, allow/deny behavior", CYAN),
    ("Networking", "Resource groups, VNets, subnets, NSGs and rule posture", ORANGE),
    ("Storage & VM", "Accounts/containers/blobs, VM state, patching, monitoring", RED),
]):
    card(slide, 0.75 + (i % 2) * 6.0, 1.75 + (i // 2) * 2.15, 5.55, 1.8, heading, body, accent=color)
text(slide, 0.9, 6.25, 11.7, 0.45, "Activity logs and security findings are derived from persisted environment state, not canned success responses.", 11.5, MUTED, True)

slide = base("KQL, Sentinel & Defender XDR", "One correlated security dataset, multiple investigation surfaces")
card(slide, 0.75, 1.7, 3.55, 2.1, "KQL engine", "Lexer → parser → AST → type checks → evaluator. Supports filtering, projection, extend, order, take, summarize, distinct, top and aggregates.", accent=CYAN)
card(slide, 4.88, 1.7, 3.55, 2.1, "Microsoft Sentinel", "Connectors, stored events, analytics rules, alerts, incidents, entities, notes and duplicate suppression.", accent=BLUE)
card(slide, 9.0, 1.7, 3.55, 2.1, "Defender XDR", "Devices, identities, process/network evidence, alerts, incidents, timelines and multi-stage correlation.", accent=ORANGE)
bullets(slide, 0.95, 4.45, 11.4, 1.5, [
    "Shared event IDs connect sign-ins, PowerShell execution, credential access, lateral movement, privilege changes and cloud activity.",
    "Detection rules execute KQL over stored rows; alerts/incidents are created from matching evidence.",
], 12)

slide = base("30-lab curriculum", "Progression from foundations to multi-stage incident investigation")
for i, (rng, heading, body, color) in enumerate([
    ("01–02", "Foundations", "Responsibility, identity & least privilege", BLUE),
    ("03–18", "AWS", "Identity → network → compute → storage hardening", CYAN),
    ("19–22", "Azure", "Entra/RBAC → network/storage → compute", BLUE),
    ("23–24", "Sentinel", "SIEM fundamentals → detections & incidents", ORANGE),
    ("25–26", "KQL", "Query fundamentals → security investigations", CYAN),
    ("27–30", "Defender XDR", "Threat detection → execution → privilege → lateral movement", RED),
]):
    y = 1.65 + i * 0.83
    pill(slide, 0.78, y, 1.1, rng, fill=SURF2, color=color)
    text(slide, 2.12, y + 0.02, 2.1, 0.3, heading, 12, FG, True)
    text(slide, 4.15, y + 0.02, 8.25, 0.42, body, 11, MUTED)

slide = base("Learning workflow", "Evidence-driven labs with three modes")
for i, (heading, body, color) in enumerate([("Guided", "Concepts + tasks + progressive hints + solution access", BLUE), ("Practice", "Hands-on workflow with policy-controlled guidance", CYAN), ("Challenge", "Concepts/hints/solution hidden; validation still required", ORANGE)]):
    card(slide, 0.75 + i * 4.05, 1.75, 3.7, 1.75, heading, body, accent=color)
for i, label in enumerate(["Baseline", "Explain risk", "Perform action", "Verify evidence", "Server validate", "Flag + achievement"]):
    x = 0.55 + i * 2.08
    pill(slide, x, 4.25, 1.75, label, fill=SURF, color=FG)
    if i < 5:
        text(slide, x + 1.77, 4.34, 0.25, 0.2, "→", 16, MUTED, True, align=PP_ALIGN.CENTER)
text(slide, 0.8, 5.35, 11.7, 0.75, "Task progress can reach only 99%. 100% completion is granted transactionally after authenticated server-side validation of persisted evidence.", 12.5, FG, True)

slide = base("Validation, flags & achievements", "Completion is a server transaction—not a UI checkbox")
for i, (stat, heading, color) in enumerate([("30/30", "validator registry", BLUE), ("Unique", "cryptographic flags", CYAN), ("Idempotent", "completion & flag reuse", ORANGE), ("7", "achievement definitions", RED)]):
    card(slide, 0.75 + i * 3.05, 1.75, 2.72, 1.6, heading, accent=color, stat=stat)
bullets(slide, 0.9, 4.0, 11.4, 1.85, [
    "Failed validation creates no completion and no flag.",
    "Two students completing the same lab receive different flags.",
    "Verified labs are immutable against reset/task/environment-state mutation paths.",
    "Hidden validator internals are not returned to the browser.",
], 12)

slide = base("Instructor, admin & learner support", "Operational workflows for classes and troubleshooting")
card(slide, 0.75, 1.7, 3.65, 2.0, "Enrollment & cohorts", "Instructor specifies seat count + expiry → one link → students choose usernames/passwords → automatic instructor assignment.", accent=ORANGE)
card(slide, 4.84, 1.7, 3.65, 2.0, "Ask a Doubt", "Per-lab chat with screen capture/upload, protected screenshot access, instructor replies and status lifecycle.", accent=CYAN)
card(slide, 8.93, 1.7, 3.65, 2.0, "Admin control center", "Users/roles/suspension, password reset, instructors, curriculum, maintenance, quotas, audit, health, performance and support.", accent=RED)
bullets(slide, 0.95, 4.45, 11.35, 1.5, [
    "Generic support tickets remain private to the reporter; administrators can reply and manage lifecycle.",
    "Doubt threads are restricted to the student, assigned instructor, and administrators.",
], 12)

slide = base("Security, performance & accessibility", "Defense-in-depth controls measured in CI")
for i, (heading, body, color) in enumerate([
    ("Web security", "HttpOnly Strict cookies, CSRF/origin checks, CSP/HSTS, clickjacking/MIME/referrer controls", BLUE),
    ("Abuse controls", "Persistent login cooldown, scoped rate limits, request/body limits, maintenance & quotas", RED),
    ("Performance", "SQLite query profiling, bounded history, performance samples/budgets", CYAN),
    ("Accessibility", "Keyboard focus, skip link, live feedback, responsive targets, reduced-motion support", ORANGE),
]):
    card(slide, 0.75 + (i % 2) * 6.0, 1.7 + (i // 2) * 2.15, 5.55, 1.8, heading, body, accent=color)

slide = base("Windows-first operations", "Local install, first admin, backup and recovery")
for i, (number, heading, body) in enumerate([
    ("1", "Download / clone", "GitHub ZIP or git clone"),
    ("2", "Run setup", "windows-setup.ps1 → npm ci + Prisma + seed"),
    ("3", "Configure Firebase", "Web config + Admin credentials + authorized domain"),
    ("4", "Provision admin", "/setup + strong one-time setup token"),
    ("5", "Start", "windows-start.ps1 → production build + server"),
    ("6", "Back up", "npm run backup before upgrades"),
]):
    card(slide, 0.75 + (i % 3) * 4.05, 1.65 + (i // 3) * 2.2, 3.7, 1.8, f"{number}. {heading}", body, accent=BLUE if i < 3 else CYAN)

slide = base("Public HTTPS with Cloudflare Tunnel", "Expose only the local web service; keep data/internal ports private")
card(slide, 0.75, 2.0, 3.25, 1.55, "CyberLab host", "Next.js on http://localhost:3000\nSQLite remains local", accent=CYAN)
text(slide, 4.05, 2.55, 0.7, 0.35, "→", 24, MUTED, True, align=PP_ALIGN.CENTER)
card(slide, 4.8, 2.0, 3.25, 1.55, "cloudflared", "Outbound tunnel\nHTTPS edge termination", accent=BLUE)
text(slide, 8.1, 2.55, 0.7, 0.35, "→", 24, MUTED, True, align=PP_ALIGN.CENTER)
card(slide, 8.85, 2.0, 3.7, 1.55, "Public hostname", "https://cyberlab.example.com\nFirebase authorized domain", accent=ORANGE)
bullets(slide, 0.95, 4.35, 11.4, 1.7, [
    "Route only the web origin; end configuration with a 404 catch-all.",
    "Do not expose SQLite, Prisma Studio, filesystem shares or internal development ports.",
    "Use Cloudflare Access as an additional perimeter layer for /admin and /setup (optionally /instructor), while keeping CyberLab role checks enabled.",
], 11.8)

slide = base("v1.0.0 release readiness", "Repository gates and operator handoff")
for i, (stat, heading, color) in enumerate([("0", "npm audit findings", CYAN), ("30", "labs with full solutions", BLUE), ("2-user", "isolation E2E", ORANGE), ("Win + Linux", "release CI targets", RED)]):
    card(slide, 0.75 + i * 3.05, 1.75, 2.72, 1.6, heading, accent=color, stat=stat)
bullets(slide, 0.9, 3.95, 11.4, 2.0, [
    "Locked npm install → audit → Prisma schema/seed → lint → TypeScript → unit/integration tests → two-user E2E → release structure check → production build.",
    "Windows setup/start scripts, SQLite backup/restore runbook, Cloudflare deployment guide and branch-protection recommendation ship with the repo.",
    "Complete DOCX lab manual and this architecture/release deck are linked from README.",
], 11.5)
text(slide, 0.9, 6.35, 11.5, 0.45, "CyberLab Studio Pro v1.0.0 — complete lab platform", 13, CYAN, True)

prs.save(out)
print(f"Generated {out}")
