import json
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

source = Path(sys.argv[1])
outdir = Path(sys.argv[2])
outdir.mkdir(parents=True, exist_ok=True)
data = json.loads(source.read_text(encoding="utf-8"))
out = outdir / "CyberLab-Studio-Pro-Lab-Manual-v1.0.0.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.55)
sec.bottom_margin = Inches(0.55)
sec.left_margin = Inches(0.62)
sec.right_margin = Inches(0.62)
styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.2)
styles["Normal"].font.color.rgb = RGBColor(29, 41, 57)
styles["Normal"].paragraph_format.space_after = Pt(3.5)
styles["Normal"].paragraph_format.line_spacing = 1.05
for name, size, color in [
    ("Title", 28, RGBColor(16, 24, 40)),
    ("Heading 1", 19, RGBColor(23, 107, 255)),
    ("Heading 2", 13, RGBColor(16, 24, 40)),
    ("Heading 3", 10.5, RGBColor(71, 84, 103)),
]:
    style = styles[name]
    style.font.name = "Aptos Display" if name in {"Title", "Heading 1"} else "Aptos"
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = True
    style.paragraph_format.space_before = Pt(8)
    style.paragraph_format.space_after = Pt(4)
if "Code Block" not in styles:
    code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
else:
    code_style = styles["Code Block"]
code_style.font.name = "Consolas"
code_style.font.size = Pt(8.2)
code_style.font.color.rgb = RGBColor(16, 24, 40)
code_style.paragraph_format.space_before = Pt(2)
code_style.paragraph_format.space_after = Pt(2)
code_style.paragraph_format.left_indent = Inches(0.1)
code_style.paragraph_format.right_indent = Inches(0.1)

for section in doc.sections:
    h = section.header.paragraphs[0]
    h.text = "CYBERLAB STUDIO PRO  |  v1.0.0 LAB MANUAL"
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in h.runs:
        r.font.size = Pt(7.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(102, 112, 133)
    f = section.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = f.add_run("CyberLab Studio Pro v1.0.0  •  Page ")
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor(102, 112, 133)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    f._p.append(fld)

def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def border_paragraph(p, fill="F2F4F7", color="D0D5DD"):
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "5")
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), color)
        p_bdr.append(el)
    p_pr.append(p_bdr)

def label_para(label, body, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label + " ")
    r.bold = True
    r.font.color.rgb = color or RGBColor(23, 107, 255)
    p.add_run(body)
    return p

def bullet(body):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(body)

def command_block(command):
    p = doc.add_paragraph(style="Code Block")
    border_paragraph(p, "F8FAFC", "D0D5DD")
    r = p.add_run(command["command"])
    r.bold = True
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("Purpose: ")
    r.bold = True
    p.add_run(command["purpose"])
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("Explanation: ")
    r.bold = True
    p.add_run(command["explanation"])

# Cover
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(65)
r = p.add_run("CYBERLAB")
r.font.size = Pt(13)
r.font.bold = True
r.font.color.rgb = RGBColor(23, 107, 255)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Studio Pro")
r.font.size = Pt(34)
r.font.bold = True
r.font.color.rgb = RGBColor(16, 24, 40)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Complete 30-Lab Manual & Solutions")
r.font.size = Pt(19)
r.font.bold = True
r.font.color.rgb = RGBColor(52, 64, 84)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Release v1.0.0")
r.font.size = Pt(12)
r.font.bold = True
r.font.color.rgb = RGBColor(8, 122, 85)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(18)
r = p.add_run("Cloud security fundamentals • AWS-compatible workflows • Azure-compatible workflows • KQL • Microsoft Sentinel • Defender XDR")
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(71, 84, 103)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(25)
r = p.add_run("Training environment boundary")
r.bold = True
r.font.color.rgb = RGBColor(23, 107, 255)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Supported cloud-style commands and investigation workflows operate on isolated CyberLab learner state. The default release does not forward learner commands to the host shell and does not contact production AWS or Azure accounts.")
doc.add_page_break()

# Intro

doc.add_heading("How to use this manual", level=1)
label_para("Audience:", "Students, instructors, and administrators using CyberLab Studio Pro v1.0.0.")
label_para("Lab modes:", "Guided and Practice can expose hints/solutions according to platform policy; Challenge hides guidance while retaining server-side validation.")
label_para("Completion:", "Task recording is learning workflow only. Final completion, flags, and achievements are issued only after server-side validation of persisted evidence.")
label_para("Account model:", "Students join through instructor-generated enrollment links and choose a username/password. Instructors are provisioned by administrators. Administrators are provisioned during first-run setup or managed by an existing administrator.")
label_para("MFA:", "Accounts can enroll Firebase TOTP two-factor authentication from Profile. Sign-in supports the real second-factor challenge flow.")
doc.add_heading("Curriculum map", level=2)
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
for i, heading in enumerate(("Area", "Labs")):
    table.rows[0].cells[i].text = heading
    shade_cell(table.rows[0].cells[i], "E8F0FF")
for area, labs in [
    ("Cloud fundamentals", "01–02"),
    ("AWS-compatible environment", "03–18"),
    ("Azure-compatible environment", "19–22"),
    ("Microsoft Sentinel", "23–24"),
    ("KQL", "25–26"),
    ("Defender XDR", "27–30"),
]:
    cells = table.add_row().cells
    cells[0].text = area
    cells[1].text = labs

doc.add_heading("All labs", level=2)
for lab in data["labs"]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(f"Lab {lab['number']:02d} — {lab['title']}")
    r.bold = True
    p.add_run(f"  |  {lab['platform'].replace('_', ' ')}  |  {lab['difficulty'].title()}  |  {lab['estimatedMinutes']} min")

for lab in data["labs"]:
    h = doc.add_heading(f"Lab {lab['number']:02d} — {lab['title']}", level=1)
    h.paragraph_format.page_break_before = True
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    for label, value in [
        ("Platform", lab["platform"].replace("_", " ")),
        ("Difficulty", lab["difficulty"].title()),
        ("Estimated time", f"{lab['estimatedMinutes']} minutes"),
    ]:
        rr = p.add_run(f"{label}: ")
        rr.bold = True
        p.add_run(value + "   ")
    label_para("Objective:", lab["objective"])
    doc.add_heading("Concept", level=2)
    doc.add_paragraph(lab["conceptExplanation"])
    doc.add_heading("Scenario", level=2)
    doc.add_paragraph(lab["scenario"])
    doc.add_heading("Prerequisites", level=2)
    for item in lab["prerequisites"]:
        bullet(item)
    doc.add_heading("Step-by-step tasks", level=2)
    for task in lab["tasks"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"Step {task['position']}: {task['title']}")
        r.bold = True
        r.font.color.rgb = RGBColor(23, 107, 255)
        p = doc.add_paragraph(task["instruction"])
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.space_after = Pt(4)
    doc.add_heading("Progressive hints", level=2)
    for index, hint in enumerate(lab["hints"], 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"Hint {index}: ")
        r.bold = True
        r.font.color.rgb = RGBColor(181, 71, 8)
        p.add_run(hint)
    solution = lab["solution"]
    doc.add_heading("Complete solution", level=2)
    label_para("Problem:", solution["problem"], RGBColor(16, 24, 40))
    label_para("Security concept:", solution["concept"], RGBColor(16, 24, 40))
    label_para("Workspace navigation:", solution["uiNavigation"], RGBColor(16, 24, 40))
    doc.add_heading("Commands / queries / actions", level=3)
    for command in solution["commands"]:
        command_block(command)
    label_para("Expected output / evidence:", solution["expectedOutput"], RGBColor(8, 122, 85))
    label_para("Why it works:", solution["whyItWorks"], RGBColor(8, 122, 85))
    label_para("Security consequence:", solution["securityConsequence"], RGBColor(180, 35, 24))
    label_para("Verification:", solution["verification"], RGBColor(23, 107, 255))
    label_para("Common mistakes:", solution["commonMistakes"], RGBColor(181, 71, 8))
    label_para("Troubleshooting:", solution["troubleshooting"], RGBColor(181, 71, 8))
    label_para("Summary:", solution["summary"], RGBColor(16, 24, 40))

doc.add_page_break()
doc.add_heading("Platform operation appendix", level=1)
doc.add_heading("Student onboarding", level=2)
for item in [
    "An instructor creates an enrollment link with an exact seat count and expiry.",
    "Each student opens the link and chooses their own username and password.",
    "The link atomically consumes one seat and assigns the account to the instructor.",
    "Public student self-registration is disabled.",
]:
    bullet(item)
doc.add_heading("Ask a Doubt", level=2)
for item in [
    "Inside a lab workspace, students can open Ask a Doubt.",
    "Students can capture a selected screen/window or upload a PNG/JPEG/WebP screenshot under 2 MB.",
    "The thread is visible only to the student, assigned instructor, and administrators.",
    "Instructor and student can continue the conversation; authorized screenshots are served through protected routes.",
]:
    bullet(item)
doc.add_heading("Public access", level=2)
doc.add_paragraph("For public HTTPS access, run CyberLab locally and publish only the web service through Cloudflare Tunnel. Keep SQLite, Prisma Studio, filesystem shares, and development/internal ports private. Cloudflare Access is recommended as an additional perimeter layer for privileged routes such as /admin and /setup.")
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(16)
r = p.add_run("Release note: ")
r.bold = True
p.add_run("This manual is generated from the same v1.0.0 curriculum definitions used by the application, so lab titles, objectives, tasks, hints, and solution steps match the shipped product.")

doc.save(out)
print(f"Generated {out}")
